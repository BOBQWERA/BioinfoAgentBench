from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "paper" / "sources" / "table_s2_source.xlsx"
OUTPUT_DIR = ROOT / "paper" / "supplementary"
OUTPUT = OUTPUT_DIR / "Table_S2.xlsx"

MEDIUM_BLACK = Side(style="medium", color="000000")
TOP_BOTTOM_BORDER = Border(top=MEDIUM_BLACK, bottom=MEDIUM_BLACK)
BOTTOM_BORDER = Border(bottom=MEDIUM_BLACK)
NO_BORDER = Border()
NO_FILL = PatternFill(fill_type=None)

HEADER_FONT = Font(name="Times New Roman", size=6.5, bold=True)
BODY_FONT = Font(name="Times New Roman", size=6.5)
ALIGN_ALL = Alignment(horizontal="left", vertical="center", wrap_text=True)


ROWS = [
    ["Planning", "P1", "Workflow Design Load", "Decomposition of a goal into an ordered, dependency-aware workflow with appropriate intermediate and final deliverables."],
    ["Planning", "P2", "Branch Decision Load", "Selection among alternative analytical routes under implicit constraints, rather than mere execution of a fixed pipeline."],
    ["Action", "A1", "Output Complexity Load", "Reliable translation of analytical intent into verifiable products that satisfy required specifications."],
    ["Action", "A2", "Methodological Constraint Load", "Sensitivity of correctness to detailed specifications and parameter choices rather than to broad directional correctness alone."],
    ["Reflection", "R1", "Semantic Repair Load", "Clarification of underdetermined semantics or incomplete problem definitions before reliable computation can proceed."],
    ["Reflection", "R2", "Data Handle Load", "Resolution of inconsistencies across data representations so that downstream analysis becomes definable and executable."],
    ["Evolution", "E1", "Experience Learning Load", "Use of abstracted prior experience as general guidance for same-class problems."],
    ["Evolution", "E2", "Strategy Transfer Load", "Adaptation of prior solution trajectories to related but non-isomorphic problems under new constraints."],
]


def main() -> None:
    source_suffix = load_source_suffix_columns()

    wb = Workbook()
    ws = wb.active
    ws.title = "Table S2"

    headers = [
        "Domain",
        "Subdomain",
        "Meaning",
        "What the probe tests",
        "Probe construction strategy",
        "Observable evidence for black-box verification",
    ]
    ws.append(headers)
    for row, suffix in zip(ROWS, source_suffix, strict=False):
        ws.append(row + suffix)

    ws.sheet_view.showGridLines = False
    ws.freeze_panes = None
    ws.auto_filter.ref = None

    widths = {
        "A": 13.0,
        "B": 10.0,
        "C": 28.0,
        "D": 52.0,
        "E": 50.0,
        "F": 52.0,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    for row in range(1, ws.max_row + 1):
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row, col)
            cell.fill = NO_FILL
            cell.alignment = ALIGN_ALL
            cell.font = HEADER_FONT if row == 1 else BODY_FONT
            if row == 1:
                cell.border = TOP_BOTTOM_BORDER
            elif row == ws.max_row:
                cell.border = BOTTOM_BORDER
            else:
                cell.border = NO_BORDER

    for row in range(1, ws.max_row + 1):
        ws.row_dimensions[row].height = estimate_row_height(ws, row)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)


def load_source_suffix_columns() -> list[list[str]]:
    wb = load_workbook(SOURCE, data_only=False)
    ws = wb["Table S2"]
    suffix = []
    current_domain = None
    for row in range(3, 11):
        domain = ws.cell(row, 1).value
        if domain is not None:
            current_domain = domain
        subdomain = ws.cell(row, 2).value
        strategy = ws.cell(row, 4).value
        evidence = ws.cell(row, 5).value
        code = str(subdomain).split(".")[0]
        suffix.append([strategy, evidence, current_domain, code])
    wb.close()

    by_code = {(code, domain): [strategy, evidence] for strategy, evidence, domain, code in suffix}
    ordered = []
    for domain, code, _, _ in ROWS:
        ordered.append(by_code[(code, domain)])
    return ordered


def estimate_row_height(ws, row: int) -> float:
    max_lines = 1
    for col in range(1, ws.max_column + 1):
        value = ws.cell(row, col).value
        if value is None:
            continue
        width = ws.column_dimensions[ws.cell(row, col).column_letter].width or 10.0
        line_width = max(1, int(width * 1.1))
        for raw_line in str(value).splitlines() or [""]:
            lines = max(1, (len(raw_line) + line_width - 1) // line_width)
            max_lines = max(max_lines, lines)
    return 11.25 if max_lines == 1 else min(11.25 + (max_lines - 1) * 7.0, 80.0)


if __name__ == "__main__":
    main()


