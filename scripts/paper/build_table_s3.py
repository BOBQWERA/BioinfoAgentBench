from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "paper" / "sources" / "frontier_case_catalog.csv"
OUTPUT_DIR = ROOT / "paper" / "supplementary"
OUTPUT = OUTPUT_DIR / "Table_S3.xlsx"

MEDIUM_BLACK = Side(style="medium", color="000000")
TOP_BOTTOM_BORDER = Border(top=MEDIUM_BLACK, bottom=MEDIUM_BLACK)
BOTTOM_BORDER = Border(bottom=MEDIUM_BLACK)
NO_BORDER = Border()
NO_FILL = PatternFill(fill_type=None)

HEADER_FONT = Font(name="Times New Roman", size=6.5, bold=True)
BODY_FONT = Font(name="Times New Roman", size=6.5)
ALIGN_ALL = Alignment(horizontal="left", vertical="center", wrap_text=True)

COLUMN_WIDTHS = {
    "title": 34.0,
    "authors": 34.0,
    "abstract": 54.0,
    "subjects": 22.0,
    "data_availability": 46.0,
    "code_availability": 40.0,
    "repository_links": 34.0,
    "doi": 18.0,
    "publication_date": 14.0,
    "url": 34.0,
    "year": 8.0,
    "journal": 16.0,
}


def main() -> None:
    df = pd.read_csv(SOURCE, encoding="utf-8-sig").fillna("")

    wb = Workbook()
    ws = wb.active
    ws.title = "Table_S3"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = None
    ws.auto_filter.ref = None

    ws.append(list(df.columns))
    for row in df.itertuples(index=False, name=None):
        ws.append([normalize_value(value) for value in row])

    apply_style(ws)
    set_widths(ws, list(df.columns))
    set_row_heights(ws)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)


def normalize_value(value):
    if pd.isna(value):
        return ""
    return value


def apply_style(ws) -> None:
    for row in range(1, ws.max_row + 1):
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row, col)
            cell.fill = NO_FILL
            cell.alignment = ALIGN_ALL
            cell.font = HEADER_FONT if row == 1 else BODY_FONT
            if row == 1:
                cell.border = TOP_BOTTOM_BORDER
            elif row == ws.max_row:
                cell.border = BOTTOM_BORDER
            else:
                cell.border = NO_BORDER


def set_widths(ws, columns: list[str]) -> None:
    for idx, name in enumerate(columns, start=1):
        width = COLUMN_WIDTHS.get(name, estimate_column_width(ws, idx))
        ws.column_dimensions[get_column_letter(idx)].width = width


def estimate_column_width(ws, col_idx: int) -> float:
    max_len = 0
    sample_limit = min(ws.max_row, 80)
    for row in range(1, sample_limit + 1):
        value = ws.cell(row, col_idx).value
        if value is None:
            continue
        for part in str(value).splitlines() or [""]:
            max_len = max(max_len, len(part))
    return min(max(8.0, max_len * 0.85 + 2.0), 36.0)


def set_row_heights(ws) -> None:
    for row in range(1, ws.max_row + 1):
        max_lines = 1
        for col in range(1, ws.max_column + 1):
            value = ws.cell(row, col).value
            if value in (None, ""):
                continue
            width = ws.column_dimensions[get_column_letter(col)].width or 10.0
            max_lines = max(max_lines, estimate_lines(str(value), width))
        ws.row_dimensions[row].height = 11.25 if max_lines == 1 else min(11.25 + (max_lines - 1) * 7.0, 220.0)


def estimate_lines(text: str, width: float) -> int:
    approx_chars_per_line = max(1, int(width * 1.1))
    total = 0
    for raw_line in text.splitlines() or [""]:
        total += max(1, (len(raw_line) + approx_chars_per_line - 1) // approx_chars_per_line)
    return total


if __name__ == "__main__":
    main()



