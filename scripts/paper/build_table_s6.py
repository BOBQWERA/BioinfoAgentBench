from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "paper" / "sources" / "frontier_tables_source.xlsx"
OUTPUT_DIR = ROOT / "paper" / "supplementary"
OUTPUT = OUTPUT_DIR / "Table_S6.xlsx"

MEDIUM_BLACK = Side(style="medium", color="000000")
TOP_BOTTOM_BORDER = Border(top=MEDIUM_BLACK, bottom=MEDIUM_BLACK)
BOTTOM_BORDER = Border(bottom=MEDIUM_BLACK)
NO_BORDER = Border()
NO_FILL = PatternFill(fill_type=None)

HEADER_FONT = Font(name="Times New Roman", size=6.5, bold=True)
BODY_FONT = Font(name="Times New Roman", size=6.5)
ALIGN_ALL = Alignment(horizontal="left", vertical="center", wrap_text=True)

SHEET_SPECS = [
    {"source": "Table S7", "target": "Pilot Rebalancing", "start_row": 3, "end_row": 5},
    {"source": "Table S7", "target": "Pilot Fingerprint", "start_row": 8, "end_row": 11},
    {"source": "Table Sx Summary", "target": "Aggregate Performance", "start_row": 2, "end_row": None},
    {"source": "Table Sy Benchmark", "target": "Benchmark Success", "start_row": 2, "end_row": None},
    {"source": "Table Sz Per-Task", "target": "Per-Task Diagnostics", "start_row": 2, "end_row": None},
    {"source": "Table Saa Subset", "target": "Subset Success", "start_row": 2, "end_row": None},
    {"source": "Table Sab High-Load", "target": "High-Load Success", "start_row": 2, "end_row": None},
    {"source": "Table Sac Replacement", "target": "Replacement Operations", "start_row": 2, "end_row": None},
    {"source": "Table Sad PARE Subdomains", "target": "PARE Rebalancing", "start_row": 2, "end_row": None},
    {"source": "Table Sae Agent Subdomains", "target": "Agent Fingerprints", "start_row": 2, "end_row": None},
    {"source": "Table Saf Task Matrix", "target": "Task Matrix", "start_row": 2, "end_row": None},
]


def main() -> None:
    source_wb = load_workbook(SOURCE, data_only=False)
    output_wb = Workbook()
    default_ws = output_wb.active
    output_wb.remove(default_ws)

    for spec in SHEET_SPECS:
        source_ws = source_wb[spec["source"]]
        output_ws = output_wb.create_sheet(spec["target"])
        copy_sheet_section(
            source_ws=source_ws,
            output_ws=output_ws,
            start_row=spec["start_row"],
            end_row=spec["end_row"] or source_ws.max_row,
        )
        style_sheet(output_ws)

    source_wb.close()
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    output_wb.save(OUTPUT)


def copy_sheet_section(source_ws, output_ws, start_row: int, end_row: int) -> None:
    merged_map = build_merged_value_map(source_ws)
    max_col = used_last_column(source_ws, merged_map, start_row, end_row)

    for source_row in range(start_row, end_row + 1):
        values = []
        for col in range(1, max_col + 1):
            values.append(get_effective_value(source_ws, merged_map, source_row, col))
        output_ws.append(values)


def build_merged_value_map(ws) -> dict[tuple[int, int], object]:
    merged_map: dict[tuple[int, int], object] = {}
    for merged_range in ws.merged_cells.ranges:
        value = ws.cell(merged_range.min_row, merged_range.min_col).value
        for row in range(merged_range.min_row, merged_range.max_row + 1):
            for col in range(merged_range.min_col, merged_range.max_col + 1):
                merged_map[(row, col)] = value
    return merged_map


def get_effective_value(ws, merged_map: dict[tuple[int, int], object], row: int, col: int):
    value = ws.cell(row, col).value
    if value is not None:
        return value
    return merged_map.get((row, col))


def used_last_column(ws, merged_map: dict[tuple[int, int], object], start_row: int, end_row: int) -> int:
    last_col = 1
    for row in range(start_row, end_row + 1):
        for col in range(1, ws.max_column + 1):
            if get_effective_value(ws, merged_map, row, col) is not None:
                last_col = max(last_col, col)
    return last_col


def style_sheet(ws) -> None:
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = None
    ws.auto_filter.ref = None

    for row in range(1, ws.max_row + 1):
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row, col)
            cell.fill = NO_FILL
            cell.alignment = ALIGN_ALL
            cell.font = HEADER_FONT if row == 1 else BODY_FONT
            if row == 1:
                cell.border = TOP_BOTTOM_BORDER
            elif row == ws.max_row:
                cell.border = BOTTOM_BORDER
            else:
                cell.border = NO_BORDER

    set_column_widths(ws)
    set_row_heights(ws)


def set_column_widths(ws) -> None:
    for col in range(1, ws.max_column + 1):
        max_len = 0
        sample_limit = min(ws.max_row, 60)
        for row in range(1, sample_limit + 1):
            value = ws.cell(row, col).value
            if value is None:
                continue
            for part in str(value).splitlines() or [""]:
                max_len = max(max_len, len(part))
        width = min(max(8.0, max_len * 0.85 + 2.0), 42.0)
        ws.column_dimensions[get_column_letter(col)].width = width


def set_row_heights(ws) -> None:
    for row in range(1, ws.max_row + 1):
        max_lines = 1
        for col in range(1, ws.max_column + 1):
            value = ws.cell(row, col).value
            if value in (None, ""):
                continue
            width = ws.column_dimensions[get_column_letter(col)].width or 10.0
            max_lines = max(max_lines, estimate_lines(str(value), width))
        if row == 1:
            height = 14.0 if max_lines == 1 else min(14.0 + (max_lines - 1) * 9.0, 220.0)
        else:
            height = 10.5 if max_lines == 1 else min(10.5 + (max_lines - 1) * 9.0, 220.0)
        ws.row_dimensions[row].height = height


def estimate_lines(text: str, width: float) -> int:
    approx_chars_per_line = max(1, int(width * 1.1))
    total = 0
    for raw_line in text.splitlines() or [""]:
        total += max(1, (len(raw_line) + approx_chars_per_line - 1) // approx_chars_per_line)
    return total


if __name__ == "__main__":
    main()


