from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "paper" / "sources" / "baseline_summary_source.xlsx"
OUTPUT_DIR = ROOT / "paper" / "supplementary"
OUTPUT = OUTPUT_DIR / "Table_S5.xlsx"

MEDIUM_BLACK = Side(style="medium", color="000000")
TOP_BOTTOM_BORDER = Border(top=MEDIUM_BLACK, bottom=MEDIUM_BLACK)
BOTTOM_BORDER = Border(bottom=MEDIUM_BLACK)
NO_BORDER = Border()
NO_FILL = PatternFill(fill_type=None)

HEADER_FONT = Font(name="Times New Roman", size=6.5, bold=True)
BODY_FONT = Font(name="Times New Roman", size=6.5)
ALIGN_ALL = Alignment(horizontal="left", vertical="center", wrap_text=True)

COLUMN_WIDTHS = {
    "A": 16.0,
    "B": 10.0,
    "C": 16.0,
    "D": 18.0,
    "E": 14.0,
    "F": 14.0,
}


def main() -> None:
    rows = load_rows()

    wb = Workbook()
    ws = wb.active
    ws.title = "Table_S5"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = None
    ws.auto_filter.ref = None

    headers = ["Benchmark", "Tasks, n", "Agent", "Successful tasks, n", "Task success, %", "95% CI, %"]
    ws.append(headers)
    for row in rows:
        ws.append(row)

    for col, width in COLUMN_WIDTHS.items():
        ws.column_dimensions[col].width = width

    apply_style(ws)
    set_row_heights(ws)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)


def load_rows() -> list[list[object]]:
    wb = load_workbook(SOURCE, data_only=False)
    ws = wb["Table S5"]

    rows = []
    current_benchmark = None
    current_tasks = None

    for row in range(3, ws.max_row + 1):
        benchmark = ws.cell(row, 1).value
        tasks = ws.cell(row, 2).value
        if benchmark is not None:
            current_benchmark = benchmark
        if tasks is not None:
            current_tasks = tasks

        rows.append(
            [
                current_benchmark,
                current_tasks,
                ws.cell(row, 3).value,
                ws.cell(row, 4).value,
                ws.cell(row, 5).value,
                ws.cell(row, 6).value,
            ]
        )

    wb.close()
    return rows


def apply_style(ws) -> None:
    for row in range(1, ws.max_row + 1):
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row, col)
            cell.fill = NO_FILL
            cell.alignment = ALIGN_ALL
            cell.font = HEADER_FONT if row == 1 else BODY_FONT
            if row == 1:
                cell.border = TOP_BOTTOM_BORDER
            elif row == ws.max_row:
                cell.border = BOTTOM_BORDER
            else:
                cell.border = NO_BORDER


def set_row_heights(ws) -> None:
    for row in range(1, ws.max_row + 1):
        max_lines = 1
        for col in range(1, ws.max_column + 1):
            value = ws.cell(row, col).value
            if value in (None, ""):
                continue
            width = ws.column_dimensions[get_column_letter(col)].width or 10.0
            max_lines = max(max_lines, estimate_lines(str(value), width))
        ws.row_dimensions[row].height = 11.25 if max_lines == 1 else min(11.25 + (max_lines - 1) * 7.0, 60.0)


def estimate_lines(text: str, width: float) -> int:
    approx_chars_per_line = max(1, int(width * 1.1))
    total = 0
    for raw_line in text.splitlines() or [""]:
        total += max(1, (len(raw_line) + approx_chars_per_line - 1) // approx_chars_per_line)
    return total


if __name__ == "__main__":
    main()



