from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "paper" / "sources" / "agent_selection_source.xlsx"
OUTPUT_DIR = ROOT / "paper" / "supplementary"
OUTPUT = OUTPUT_DIR / "Table_S4.xlsx"

MEDIUM_BLACK = Side(style="medium", color="000000")
TOP_BOTTOM_BORDER = Border(top=MEDIUM_BLACK, bottom=MEDIUM_BLACK)
BOTTOM_BORDER = Border(bottom=MEDIUM_BLACK)
NO_BORDER = Border()
NO_FILL = PatternFill(fill_type=None)

HEADER_FONT = Font(name="Times New Roman", size=6.5, bold=True)
BODY_FONT = Font(name="Times New Roman", size=6.5)
ALIGN_ALL = Alignment(horizontal="left", vertical="center", wrap_text=True)

HEADERS = [
    "System",
    "System Type",
    "Year",
    "Venue or Status",
    "Public Availability",
    "Primary Capability Emphasis",
    "Relevance to Bioinformatics Workflows",
    "Rationale for Inclusion",
]

COLUMN_WIDTHS = {
    "A": 16.0,
    "B": 28.0,
    "C": 8.0,
    "D": 20.0,
    "E": 14.0,
    "F": 40.0,
    "G": 40.0,
    "H": 38.0,
}


def main() -> None:
    rows = load_integrated_rows()

    wb = Workbook()
    ws = wb.active
    ws.title = "Table_S4"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = None
    ws.auto_filter.ref = None

    ws.append(HEADERS)
    for row in rows:
        ws.append(row)

    for col, width in COLUMN_WIDTHS.items():
        ws.column_dimensions[col].width = width

    apply_style(ws)
    set_row_heights(ws)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)


def load_integrated_rows() -> list[list[object]]:
    wb = load_workbook(SOURCE, data_only=False)
    table5 = wb["Table 5"]
    table_s4 = wb["Table S4"]

    base_rows = []
    for row in range(3, table5.max_row + 1):
        system = table5.cell(row, 1).value
        if not system:
            continue
        base_rows.append(
            {
                "System": system,
                "System Type": table5.cell(row, 2).value,
                "Year": table5.cell(row, 3).value,
                "Venue or Status": table5.cell(row, 4).value,
                "Public Availability": table5.cell(row, 5).value,
            }
        )

    detail_rows = {}
    for row in range(3, table_s4.max_row + 1):
        system = table_s4.cell(row, 1).value
        if not system:
            continue
        detail_rows[system] = {
            "Primary Capability Emphasis": table_s4.cell(row, 3).value,
            "Relevance to Bioinformatics Workflows": table_s4.cell(row, 4).value,
            "Rationale for Inclusion": table_s4.cell(row, 5).value,
        }

    wb.close()

    merged = []
    for entry in base_rows:
        system = entry["System"]
        details = detail_rows.get(system)
        if details is None:
            raise RuntimeError(f"Missing Table S4 details for system: {system}")
        merged.append(
            [
                entry["System"],
                entry["System Type"],
                entry["Year"],
                entry["Venue or Status"],
                entry["Public Availability"],
                details["Primary Capability Emphasis"],
                details["Relevance to Bioinformatics Workflows"],
                details["Rationale for Inclusion"],
            ]
        )
    return merged


def apply_style(ws) -> None:
    for row in range(1, ws.max_row + 1):
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row, col)
            cell.fill = NO_FILL
            cell.alignment = ALIGN_ALL
            cell.font = HEADER_FONT if row == 1 else BODY_FONT
            if row == 1:
                cell.border = TOP_BOTTOM_BORDER
            elif row == ws.max_row:
                cell.border = BOTTOM_BORDER
            else:
                cell.border = NO_BORDER


def set_row_heights(ws) -> None:
    for row in range(1, ws.max_row + 1):
        max_lines = 1
        for col in range(1, ws.max_column + 1):
            value = ws.cell(row, col).value
            if value in (None, ""):
                continue
            width = ws.column_dimensions[get_column_letter(col)].width or 10.0
            max_lines = max(max_lines, estimate_lines(str(value), width))
        ws.row_dimensions[row].height = 11.25 if max_lines == 1 else min(11.25 + (max_lines - 1) * 7.0, 120.0)


def estimate_lines(text: str, width: float) -> int:
    approx_chars_per_line = max(1, int(width * 1.1))
    total = 0
    for raw_line in text.splitlines() or [""]:
        total += max(1, (len(raw_line) + approx_chars_per_line - 1) // approx_chars_per_line)
    return total


if __name__ == "__main__":
    main()



